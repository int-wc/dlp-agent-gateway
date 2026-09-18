"""Bounded extraction for the isolated analyzer worker."""

import io
import re
import zipfile

from docx import Document
from pypdf import PdfReader


MAX_BYTES = 8 * 1024 * 1024
MAX_TEXT = 100_000
TEXT_SUFFIXES = {
    ".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".go",
    ".java", ".yaml", ".yml", ".log",
}


class ParseFailure(Exception):
    """The document could not be completely and safely extracted."""


def _bounded(text: str) -> str:
    if len(text) > MAX_TEXT:
        raise ParseFailure()
    return text


def extract(data: bytes, filename: str) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf") and data.startswith(b"%PDF-"):
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted or len(reader.pages) > 30:
                raise ParseFailure()
            return _bounded("\n".join((page.extract_text() or "") for page in reader.pages))
        if name.endswith(".docx") and zipfile.is_zipfile(io.BytesIO(data)):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 500 or sum(entry.file_size for entry in entries) > 16 * 1024 * 1024:
                    raise ParseFailure()
                if any(entry.file_size > 100 * max(entry.compress_size, 1) for entry in entries):
                    raise ParseFailure()
                if "[Content_Types].xml" not in archive.namelist():
                    raise ParseFailure()
            document = Document(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
            return _bounded("\n".join(paragraphs))
        if (name.endswith(".png") and data.startswith(b"\x89PNG\r\n\x1a\n")) or (
            name.endswith((".jpg", ".jpeg")) and data.startswith(b"\xff\xd8\xff")
        ):
            try:
                from PIL import Image
                import pytesseract
            except ImportError as exc:
                raise ParseFailure() from exc
            image = Image.open(io.BytesIO(data))
            if image.width * image.height > 10_000_000:
                raise ParseFailure()
            return _bounded(pytesseract.image_to_string(image, timeout=5))
        if any(name.endswith(ext) for ext in TEXT_SUFFIXES) and b"\x00" not in data:
            return _bounded(data.decode("utf-8-sig"))
    except (ValueError, UnicodeError, OSError, RuntimeError, zipfile.BadZipFile, KeyError, ImportError) as exc:
        raise ParseFailure() from exc
    raise ParseFailure()


def signals(text: str) -> list[str]:
    patterns = {
        "private_key": r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----",
        "aws_access_key": r"\bAKIA[0-9A-Z]{16}\b",
        "chinese_id_candidate": r"(?<!\d)\d{17}[0-9Xx](?!\d)",
        "phone_candidate": r"(?<!\d)1[3-9]\d{9}(?!\d)",
    }
    return [label for label, pattern in patterns.items() if re.search(pattern, text)]
